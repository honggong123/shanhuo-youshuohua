# -*- coding: utf-8 -*-
"""生成《山货有话说》竞赛作品说明书 docx。

用法：python generate_shuomingshu.py
输出：作品说明书-山货有话说.docx（同目录）
"""
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT = "作品说明书-山货有话说.docx"

HEI = "黑体"
SONG = "宋体"
KAI = "楷体"


def set_font(run, name_cn=SONG, name_en="Times New Roman", size=12, bold=False, color="000000"):
    run.font.name = name_en
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name_cn)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def body_para(doc, text, indent=True, size=12, bold=False, align=None, space_after=6):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = 1.5
    pf.space_after = Pt(space_after)
    if indent:
        pf.first_line_indent = Pt(size * 2)
    if align is not None:
        p.alignment = align
    run = p.add_run(text)
    set_font(run, SONG, size=size, bold=bold)
    return p


def heading(doc, text, level=1):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(14 if level == 1 else 10)
    pf.space_after = Pt(8)
    pf.keep_with_next = True
    run = p.add_run(text)
    if level == 1:
        set_font(run, HEI, size=16, bold=True)
    else:
        set_font(run, HEI, size=13, bold=True)
    # 使用 Word 内置标题样式层级，便于导航窗格识别
    p.style = doc.styles[f"Heading {level}"]
    # 样式覆盖后重新设置字体（内置样式自带蓝色 Calibri）
    for r in p.runs:
        set_font(r, HEI, size=16 if level == 1 else 13, bold=True, color="000000")
    return p


def info_table(doc, rows):
    table = doc.add_table(rows=len(rows), cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (k, v) in enumerate(rows):
        c1, c2 = table.rows[i].cells
        c1.width = Cm(4.2)
        c2.width = Cm(11.0)
        r1 = c1.paragraphs[0].add_run(k)
        set_font(r1, HEI, size=11, bold=True)
        r2 = c2.paragraphs[0].add_run(v)
        set_font(r2, SONG, size=11)
        for c in (c1, c2):
            c.paragraphs[0].paragraph_format.space_before = Pt(3)
            c.paragraphs[0].paragraph_format.space_after = Pt(3)
    return table


doc = Document()

# 页面设置：A4，常规页边距
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
sec.top_margin = sec.bottom_margin = Cm(2.54)
sec.left_margin = sec.right_margin = Cm(2.8)

# ============ 标题页 ============
for _ in range(5):
    doc.add_paragraph()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_font(p.add_run("《山货有话说》"), HEI, size=26, bold=True)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_font(p.add_run("—— AI 助农数字文创平台"), HEI, size=18, bold=True)
doc.add_paragraph()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_font(p.add_run("作 品 说 明 书"), KAI, size=16)
for _ in range(3):
    doc.add_paragraph()

info_table(doc, [
    ("作品名称", "山货有话说——AI 助农数字文创平台"),
    ("参赛赛道", "指定命题赛道（乡村振兴·科创赋能 / AIGC 类数字创意作品）"),
    ("作品类别", "移动应用开发 / 人工智能应用"),
    ("关键词", "乡村振兴；AIGC；多模态大模型；数字文创；智能语音"),
    ("团队名称", "【请填写：团队名称】"),
    ("团队成员", "【请填写：成员姓名及分工】"),
    ("指导教师", "【请填写：指导教师姓名】"),
    ("完成日期", "【请填写：____年____月____日】"),
])

doc.add_page_break()

# ============ 正文 ============
heading(doc, "一、作品概述", 1)
body_para(doc,
    "《山货有话说》是一个面向乡村振兴场景的 AI 助农数字文创平台。用户仅需拍摄或上传一张农产品照片，"
    "平台即可在 30 秒内自动完成产品识别建档，并一键生成四件营销素材：文创海报、电商卖点文案、语音介绍与短视频分镜脚本；"
    "同时内置 AI 数字讲解员，可为乡村文旅景点生成沉浸式导游词并进行语音讲解与实时问答。"
    "作品以响应式 H5 网页形式交付，手机、电脑浏览器均可直接使用，无需安装。")
body_para(doc,
    "作品对应全国大学生数字媒体科技作品及创意竞赛指定命题方向，将多模态大模型、AIGC 图像生成、"
    "大语言模型内容创作与智能语音合成四类人工智能能力，落地于助农营销与乡村文旅两大真实场景。")

heading(doc, "二、创作背景与需求分析", 1)
heading(doc, "（一）现实痛点", 2)
body_para(doc,
    "调研发现，山区优质农产品“卖难”的症结往往不在品质，而在内容与品牌能力的缺失："
    "其一，农户普遍缺乏拍摄、文案、美工与短视频制作能力，好货没有好包装；"
    "其二，专业营销团队服务费用高，小农户与合作社难以负担；"
    "其三，乡村文旅资源丰富，但讲解人才稀缺，游客“看得到风景、听不到故事”。")
heading(doc, "（二）需求提炼", 2)
body_para(doc,
    "由此提炼出两大核心需求：一是“低门槛内容生产”——让不懂营销的农户，用一部手机就能把山货讲出故事、做出海报；"
    "二是“可复制的文旅讲解”——让每一个村落都拥有一位随叫随到、能说会道的数字讲解员。")

heading(doc, "三、作品设计", 1)
heading(doc, "（一）总体架构", 2)
body_para(doc,
    "作品采用前后端分离架构。前端基于 Vue 3 与 Vite 构建移动端优先的响应式界面；"
    "后端基于 Python FastAPI 提供 RESTful 接口，统一封装 OpenAI 兼容协议对接大模型能力，"
    "本地使用 Pillow 完成海报排版合成，使用 Edge-TTS 完成语音合成。整体链路轻量，普通服务器或笔记本即可部署运行。")
heading(doc, "（二）核心功能模块", 2)
body_para(doc, "1. 山货四件套生成器：用户上传农产品照片后，多模态大模型自动识别产品名称、品类、产地与品质亮点，形成数字档案；"
    "随后并行调用三条生成链路——大语言模型撰写卖点文案与分镜脚本、图像生成模型绘制国潮风海报底图并叠加精确中文排版、"
    "智能语音将介绍文本转换为自然流畅的语音，最终“海报、文案、语音、脚本”四件套一次成型。", bold=False)
body_para(doc, "2. AI 数字讲解员：内置村落知识库，可为选定村落生成富有画面感的沉浸式导游词，自动配音讲解；"
    "游客亦可通过自然语言提问，获得口语化的即时回答。")
body_para(doc, "3. 演示兜底机制：未配置模型密钥时，平台自动进入演示模式，以内置精修数据与本地模板海报保证全流程可用，"
    "既便于开发调试，也保障评审现场弱网环境下的稳定演示。")

heading(doc, "四、关键技术与实现", 1)
body_para(doc, "（1）多模态识图：采用视觉—语言大模型对农产品照片进行结构化识别，通过提示词工程约束输出 JSON 格式，保证下游环节稳定消费。")
body_para(doc, "（2）统一模型网关：后端以 OpenAI 兼容接口统一封装文本、视觉、图像生成三类模型，"
    "更换模型厂商仅需修改环境变量，实现“一套代码、多厂可插拔”。")
body_para(doc, "（3）中文海报精确排版：针对文生图模型中文渲染易乱码的行业痛点，创新采用“AI 底图 + 本地精确文字排版”两段式合成方案——"
    "先由图像模型生成无文字的国潮风底图，再由本地排版引擎叠加标题、标语与产地徽章，保证海报中文始终清晰准确。")
body_para(doc, "（4）智能语音：基于 Edge-TTS 实现高质量中文语音合成，零密钥、零成本。")
body_para(doc, "（5）前后端工程化：前端 Vue 3 + Vite + Tailwind CSS 构建响应式界面，构建产物由后端统一托管，实现“一条命令启动完整作品”。")

heading(doc, "五、创新点分析", 1)
body_para(doc, "创新点一：单张照片驱动的全链路内容生产。以一张照片为唯一输入，串联识别、文案、图像、语音、脚本五类能力，"
    "把专业营销团队数天的工作压缩到 30 秒内完成。", indent=True)
body_para(doc, "创新点二：“助农 + 文旅”双场景闭环。同一平台既解决“卖难”又解决“游难”，产品与场景互相导流，叙事完整。")
body_para(doc, "创新点三：AI 底图与本地排版相结合的中文海报合成方案，兼顾生成式美术的表现力与版式文字的准确性。")
body_para(doc, "创新点四：零门槛可用性设计。演示模式使作品在不联网、无密钥的环境下依然完整可演示，显著降低使用与评审门槛。")

heading(doc, "六、应用价值与社会意义", 1)
body_para(doc, "对农户与合作社：免费、低门槛地获得专业级营销素材，降低农产品上行门槛，直接助力增收。")
body_para(doc, "对县域文旅：以数字讲解员补齐乡村讲解人才缺口，提升游客体验，带动乡村旅游二次消费。")
body_para(doc, "对时代命题：作品将人工智能技术切实应用于乡村振兴一线，是“科技赋能三农”的具体实践，"
    "呼应国家数字乡村战略，具有清晰的社会价值与推广意义。")

heading(doc, "七、应用前景与推广计划", 1)
body_para(doc, "短期计划：与本地 1~2 家农产品合作社及村集体开展试用合作，以真实素材持续调优提示词与生成质量。"
    "中期计划：拓展微信小程序端与方言语音包，接入产销数据可视化看板，形成“内容 + 数据”双轮驱动。"
    "长期愿景：沉淀为县域级数字助农基础设施，向更多革命老区与脱贫地区推广。")

heading(doc, "八、作品演示与使用说明", 1)
body_para(doc, "作品为响应式 Web 应用，按项目 README 执行两条命令即可启动：后端目录下执行 pip install -r requirements.txt 安装依赖，"
    "再执行 python run.py 启动服务，浏览器访问 http://127.0.0.1:8000 即可体验完整功能。"
    "在 backend/.env 中填入任意一家大模型服务的 API Key（推荐智谱 GLM 免费模型），即可从演示模式无缝切换为真实 AI 生成。")

heading(doc, "九、总结", 1)
body_para(doc,
    "《山货有话说》用一张照片打通了山货从“被看见”到“被记住”的完整链路。它不是一项炫技的技术堆砌，"
    "而是一次面向真实乡土需求的 AI 落地实践。我们相信，当技术愿意俯下身来讲述土地的故事，乡村振兴就有了更温柔也更强大的助力。")

doc.save(OUT)
print("已生成:", OUT)
